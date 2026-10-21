import re, collections, textwrap
from pathlib import Path
from io import BytesIO
from xml.sax.saxutils import escape
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, PageBreak,
    NextPageTemplate, Image as RLImage, KeepTogether)
from reportlab.platypus.tableofcontents import TableOfContents
from _editorial_data import CHAPTER_TITLES, CHAPTER_START_PAGES, FIGURES

SOURCE='/tmp/arena-coding-uploads/CPY_902_Advanced_Psychopathology_Textbook.pdf'
OUTDIR=Path('/home/user/tbugeusa/deliverables');OUTDIR.mkdir(parents=True,exist_ok=True)
PDF_OUT=OUTDIR/'CPY_902_Advanced_Psychopathology_Final.pdf'
DOCX_OUT=OUTDIR/'CPY_902_Advanced_Psychopathology_Final.docx'
COVER_SOURCE=OUTDIR/'assets'/'CPY_902_Advanced_Psychopathology_Front_Cover.jpeg'
FONTDIR='/tmp/source-serif-font/TTF'
FONT_REG=FONTDIR+'/SourceSerif4-Regular.ttf';FONT_BOLD=FONTDIR+'/SourceSerif4-Bold.ttf'
FONT_ITALIC=FONTDIR+'/SourceSerif4-It.ttf';FONT_BOLDITALIC=FONTDIR+'/SourceSerif4-BoldIt.ttf'
SANS='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf';SANS_BOLD='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'

GLOBAL_REMOVE_PATTERNS=[
 re.compile(r'Reasoning of this kind helps counsellors, psychologists, physicians, families, and clients discuss .*? using a shared account of evidence and uncertainty\.',re.I),
 re.compile(r'This method reduces premature closure because the reasoning about .*? can be checked against alternatives such as .*? and documented for later review\.',re.I),
]
PER_CHAPTER_KEEP_PATTERNS={
 'provisional':re.compile(r'The formulation remains provisional; new evidence about .*? may require revision in the discussion of .*?\.',re.I),
 'good_practice':re.compile(r'Good practice combines diagnostic knowledge with collaboration and repeated reassessment, especially when .*? changes across time or setting\.',re.I),
 'practical_way':re.compile(r'A practical way to study .*? is to move from description to hypothesis and then back to evidence\.',re.I),
 'immediate_aim':re.compile(r'The immediate aim is an explanation that can guide .*? while making uncertainty about .*? explicit\.',re.I),
 'symptom_meaning':re.compile(r'(?:At the same time, |In many cases, |From a formulation perspective, |In practice, |Longitudinally, |Across the lifespan, )?a symptom can carry different meanings across developmental stages and cultural settings\.',re.I),
 'records':re.compile(r'(?:In many cases, |Where risk is present, |In service settings, |At the same time, |For assessment purposes, |Clinically, |Across the lifespan, |From a developmental perspective, )?records should distinguish the client’s report, observed behaviour, collateral information, test results, interpretation, and planned follow-up\.',re.I),
 'strengths':re.compile(r'(?:In many cases, |Where risk is present, |At the same time, |For assessment purposes, |When the history is incomplete, |In practice, |Across the lifespan, |Longitudinally, |Clinically, |From a developmental perspective, )?the person’s strengths, routines, supportive relationships, and previous periods of recovery offer evidence about protective processes\.',re.I),
 'culture':re.compile(r'(?:Where risk is present, |At the same time, |In many cases, |For assessment purposes, |Across the lifespan, |Clinically, |In practice, |Longitudinally, |From a developmental perspective, )?cultural and family explanations should be explored directly rather than inferred from group membership\.',re.I),
 'measurement':re.compile(r'(?:In many cases, |In practice, |Across the lifespan, |In service settings, |From a formulation perspective, |Clinically, |Longitudinally, )?measurement results should be interpreted against language, literacy, developmental level, cultural context, response style, and the purpose for which the measure was designed\.',re.I),
 'counsellor_referral':re.compile(r'(?:In practice, |In service settings, |Where risk is present, |Across the lifespan, |From a developmental perspective, |Clinically, |Longitudinally, |From a formulation perspective, )?a counsellor should recognise when the presentation requires medical assessment, prescribing, nutritional rehabilitation, neurological investigation, emergency intervention, or another specialist service\.',re.I),
 'context_data':re.compile(r'Context is part of the clinical data for .*? because the meaning and impact of .*? can change across settings\. Behaviour that looks unusual in one setting may be understandable in another; for .*?, the assessor asks whether the response is protective, restrictive, or both at different times\.',re.I),
 'environment_demands':re.compile(r'(?:In|During|At) [^,]{1,90}, practitioners should ask what demands the environment places on attention, emotion, social communication, sleep, decision-making, and self-regulation\.',re.I),
 'assessment_observations':re.compile(r'An assessment may identify .*?, but these observations become clinically useful only after they are linked to timing, severity, triggers, consequences, and recovery periods\.',re.I),
 'premature_selection':re.compile(r'Comparing .*? with .*? reduces the risk of selecting a familiar label first and then bending later information to confirm it\.',re.I),
 'respectful_curiosity':re.compile(r'Respectful curiosity about .*?, language, and context can improve diagnostic accuracy while supporting engagement\.',re.I),
 'pattern_interactions':re.compile(r'The pattern may involve .*?, while contextual pressures interact with .*?\.',re.I),
}
GLOBAL_KEEP_FIRST_PATTERNS={
 'central_task':re.compile(r'The central task is to connect phenomenology, impairment, context, and course without reducing the person to a diagnosis\.',re.I),
 'area_patterns':re.compile(r'This area is best understood through patterns that unfold across time, setting, development, and relationships\.',re.I),
 'discipline':re.compile(r'The most important discipline is repeated testing of the working formulation against chronology, impairment, context, physical health, culture, development, and response to care\.',re.I),
 'diagnostic_language':re.compile(r'Diagnostic language should support communication and treatment selection without erasing the person’s strengths, preferences, and social circumstances\.',re.I),
}

def replace_case(text,old,new):
 pat=re.compile(r'\b'+re.escape(old)+r'\b',re.I)
 def f(m):
  s=m.group(0)
  if s.isupper():return new.upper()
  if s[0].isupper():return new[0].upper()+new[1:]
  return new
 return pat.sub(f,text)

def british_english(text):
 protected={'PhD in Counseling Psychology':'__OFFICIAL_PROGRAMME_TITLE__','Centers for Disease Control and Prevention':'__CDC_NAME__','World Health Organization':'__WHO_NAME__','National Institute of Mental Health':'__NIMH_NAME__','Substance Abuse and Mental Health Services Administration':'__SAMHSA_NAME__'}
 for k,v in protected.items():text=text.replace(k,v)
 mapping=[('behavioral','behavioural'),('behaviorally','behaviourally'),('behavior','behaviour'),('counseling','counselling'),('counselor','counsellor'),('counselors','counsellors'),('programs','programmes'),('program','programme'),('analyze','analyse'),('analyzed','analysed'),('analyzing','analysing'),('organize','organise'),('organized','organised'),('organizing','organising'),('organization','organisation'),('organizations','organisations'),('recognize','recognise'),('recognized','recognised'),('recognizing','recognising'),('standardize','standardise'),('standardized','standardised'),('standardizing','standardising'),('prioritize','prioritise'),('prioritized','prioritised'),('prioritizing','prioritising'),('minimize','minimise'),('minimized','minimised'),('minimizing','minimising'),('emphasize','emphasise'),('emphasized','emphasised'),('emphasizing','emphasising'),('stigmatize','stigmatise'),('stigmatized','stigmatised'),('stigmatizing','stigmatising'),('modeling','modelling'),('modeled','modelled'),('labeling','labelling'),('labeled','labelled'),('traveling','travelling'),('traveled','travelled'),('canceled','cancelled'),('canceling','cancelling'),('judgment','judgement'),('etiology','aetiology'),('etiological','aetiological'),('pediatric','paediatric'),('aging','ageing'),('defense','defence'),('anemia','anaemia'),('centered','centred'),('centering','centring'),('centers','centres'),('center','centre')]
 for a,b in mapping:text=replace_case(text,a,b)
 for k,v in protected.items():text=text.replace(v,k)
 text=re.sub(r'\bDSM\s*[-–—]?\s*5\s*[-–—]?\s*TR\b','DSM-5-TR',text,flags=re.I)
 text=re.sub(r'\bICD\s*[-–—]?\s*11\b','ICD-11',text,flags=re.I)
 for abbr in ['ADHD','OCD','PTSD','ERP','CBT','DBT','USA','PhD']:text=re.sub(r'\b'+abbr+r'\b',abbr,text,flags=re.I)
 text=text.replace('NOT-FOR-SALE','NOT FOR SALE').replace('NOT -FOR-SALE','NOT FOR SALE')
 text=re.sub(r'\s+',' ',text).strip()
 if text and text[0].islower():text=text[0].upper()+text[1:]
 text=re.sub(r'\bstandardised measures differs\b','standardised measures differ',text,flags=re.I)
 text=re.sub(r'\bmeasurement results differs\b','measurement results differ',text,flags=re.I)
 text=re.sub(r'\bmeasures differs\b','measures differ',text,flags=re.I)
 text=re.sub(r'\ba integrated\b','an integrated',text,flags=re.I)
 text=re.sub(r'\s+([,.;:!?])',r'\1',text)
 text=re.sub(r'\(\s+','(',text);text=re.sub(r'\s+\)',')',text)
 text=re.sub(r'\bsymptom symptom clusters\b','symptom clusters',text,flags=re.I)
 text=re.sub(r'\ba operations supervisor\b','an operations supervisor',text,flags=re.I)
 plural_subjects=r'(difficulties|patterns|problems|symptoms|thoughts|behaviours|attacks|compulsions|obsessions|features|factors|conditions|disorders|effects|processes|responses|concerns|episodes|changes|experiences|events|signs|risks|differences|beliefs|rates|biases|relationships|contributors|interests|rituals|instruments|interviews|ratings|scales|measures|results|findings|observations|reports|records|resources|services|needs|mechanisms|pathways|influences|outcomes|questions|issues|cases)'
 text=re.sub(r'\b'+plural_subjects+r' is (present|accepted|explored)\b',r'\1 are \2',text,flags=re.I)
 text=re.sub(r'\b(screening instruments|structured and semi-structured interviews|behaviour ratings|severity scales) differs\b',r'\1 differ',text,flags=re.I)
 text=re.sub(r'\b(features|rituals|interests|difficulties|concerns|effects) varies across settings\b',r'\1 vary across settings',text,flags=re.I)
 text=re.sub(r'\b(relationships|rates|effects|differences|beliefs|experiences|conditions|contributors|factors|patterns) has altered\b',r'\1 have altered',text,flags=re.I)
 text=re.sub(r'(?<=[.!?])\s+([a-z])',lambda m:' '+m.group(1).upper(),text)
 return text

def remove_repetitive_boilerplate(text,ch,seen_by_chapter,seen_global):
 local=seen_by_chapter[ch]
 for pat in GLOBAL_REMOVE_PATTERNS:text=pat.sub('',text)
 for key,pat in GLOBAL_KEEP_FIRST_PATTERNS.items():
  def keep_global(m):
   if key in seen_global:return ''
   seen_global.add(key);return m.group(0)
  text=pat.sub(keep_global,text)
 for key,pat in PER_CHAPTER_KEEP_PATTERNS.items():
  def keep_local(m):
   if key in local['categories']:return ''
   local['categories'].add(key);return m.group(0)
  text=pat.sub(keep_local,text)
 parts=re.split(r'(?<=[.!?])\s+(?=[A-Z])',text);out=[]
 for sentence in parts:
  s=sentence.strip()
  if not s:continue
  key=re.sub(r'[^a-z0-9]+',' ',s.lower()).strip()
  if len(key.split())>=12 and key in local['sentences']:continue
  if len(key.split())>=12:local['sentences'].add(key)
  out.append(s)
 text=' '.join(out)
 text=re.sub(r'\s+([,.;:!?])',r'\1',text);text=re.sub(r'\s{2,}',' ',text).strip()
 return text

def join_lines(lines):
 s=''
 for line in lines:
  line=line.strip()
  if not line:continue
  if not s:s=line
  elif s.endswith('-') and line[:1].islower():s+=line
  else:s+=' '+line
 return re.sub(r'\s+',' ',s).strip()

def group_visual_lines(lines):
 grouped=[]
 for y0,x0,font,size,txt in sorted(lines,key=lambda x:(x[0],x[1])):
  if grouped and abs(y0-grouped[-1][0])<=2 and font==grouped[-1][2] and abs(size-grouped[-1][3])<.1:
   gy,gx,gf,gs,gt=grouped[-1];sep='' if gt.endswith('-') else ' ';grouped[-1]=(gy,gx,gf,gs,(gt+sep+txt).strip())
  else:grouped.append((y0,x0,font,size,txt))
 return grouped

pdf=pymupdf.open(SOURCE)
chapters={i:{'title':CHAPTER_TITLES[i-1],'objectives':[],'sections':collections.OrderedDict(),'source_pages':[],'source_word_count':None} for i in range(1,17)}
current_ch=0;current_section=None;current_lines=[]
def flush_body():
 global current_lines,current_ch,current_section
 if current_lines and current_ch in chapters and current_section:
  para=join_lines(current_lines)
  if para:
   if current_section=='__OBJECTIVES__':chapters[current_ch]['objectives'].append(para)
   else:chapters[current_ch]['sections'].setdefault(current_section,[]).append(para)
 current_lines=[]
for pnum,page in enumerate(pdf,1):
 lines=[]
 for block in page.get_text('dict')['blocks']:
  if 'lines' not in block:continue
  for line in block['lines']:
   spans=line['spans'];txt=''.join(s['text'] for s in spans).strip()
   if not txt:continue
   x0=min(s['bbox'][0] for s in spans);y0=min(s['bbox'][1] for s in spans);font=spans[0]['font'];size=spans[0]['size']
   if y0<60 or y0>780:continue
   lines.append((y0,x0,font,size,txt))
 lines=group_visual_lines(lines)
 for y0,x0,font,size,txt in lines:
  if txt=='Selected References and Current Diagnostic Resources':
   flush_body();current_ch=0;current_section=None;continue
  if 'Calibri-Bold' in font and size>=15.5 and txt.upper().startswith('CHAPTER'):
   flush_body();m=re.search(r'CHAPTER\s+(\d+)',txt.upper())
   if m:current_ch=int(m.group(1));current_section='__TITLE__';chapters[current_ch]['source_pages'].append(pnum)
   continue
  if current_ch and 'Calibri-Bold' in font and size>=15.5 and current_section=='__TITLE__':continue
  if current_ch and 'Calibri-Bold' in font and 13.5<=size<15.5:
   flush_body();current_section=txt;chapters[current_ch]['sections'].setdefault(current_section,[]);continue
  if current_ch and font.startswith('TimesNewRomanPSMT') and 11.5<=size<=12.1:
   if current_section is None or current_section=='__TITLE__':current_section='__OBJECTIVES__'
   if txt.lower().startswith('verified chapter body word count:'):
    flush_body();m=re.search(r':\s*([\d,]+)\s+words',txt,re.I)
    if m:chapters[current_ch]['source_word_count']=int(m.group(1).replace(',',''))
    continue
   if current_lines and x0>95:flush_body()
   current_lines.append(txt)
  elif current_lines and current_section!='__TITLE__':flush_body()
flush_body()

# The original front matter has no missing text beyond the explicit author placeholder.
front={'Publication and Academic Use Notice':[],'Preface':[]};cur_front=None;front_lines=[]
def flush_front():
 global front_lines,cur_front
 if front_lines and cur_front in front:
  p=join_lines(front_lines)
  if p:front[cur_front].append(p)
 front_lines=[]
for pnum in [2,3,4]:
 page=pdf[pnum-1];lines=[]
 for block in page.get_text('dict')['blocks']:
  if 'lines' not in block:continue
  for line in block['lines']:
   spans=line['spans'];txt=''.join(s['text'] for s in spans).strip()
   if not txt:continue
   x0=min(s['bbox'][0] for s in spans);y0=min(s['bbox'][1] for s in spans);font=spans[0]['font'];size=spans[0]['size']
   if y0<60 or y0>780:continue
   lines.append((y0,x0,font,size,txt))
 lines=group_visual_lines(lines)
 for y0,x0,font,size,txt in lines:
  if 'Calibri-Bold' in font and size>=15.5:
   flush_front();cur_front=txt if txt in front else None;continue
  if cur_front and font.startswith('TimesNewRomanPSMT') and 11.5<=size<=12.1:
   if front_lines and x0>95:flush_front()
   front_lines.append(txt)
flush_front()
notice=front['Publication and Academic Use Notice']
preface=front['Preface']
if len(notice)!=4 or len(preface)!=4:raise ValueError(f'Unexpected source front-matter paragraph counts: notice={len(notice)}, preface={len(preface)}')

def parse_objectives(raw):
 if not raw:return []
 t=raw[0];t=re.sub(r'^Learning objectives\.\s*','',t,flags=re.I)
 t=re.sub(r'^By the end of this chapter, the doctoral student should be able to\s*','',t,flags=re.I)
 parts=[re.sub(r'^and\s+','',x.strip(' .;'),flags=re.I) for x in t.split(';') if x.strip(' .;')]
 return [british_english(x[0].upper()+x[1:]+('.' if not x.endswith(('.', '?', '!')) else '')) for x in parts]

EXPECTED_SECTION_HEADINGS=[
 'Chapter Overview','Orientation and Scope','Phenomenology and Clinical Presentation',
 'Historical Development of the Concept','DSM-5-TR and ICD-11 Perspectives',
 'Dimensional and Transdiagnostic Thinking','Development and Course',
 'Genetic and Neurobiological Contributions','Learning, Cognition, Emotion, and Behaviour',
 'Relationships, Family, and Social Context','Culture, Language, Sex, Gender, and Identity',
 'Assessment Strategy','Differential Diagnosis','Comorbidity and Diagnostic Overlap',
 'Functional Impact and Quality of Life','Risk, Safety, and Protective Factors',
 'Case Formulation','Treatment and Interdisciplinary Care','Real-Life Applications',
 'Ethics, Stigma, and Professional Communication','Current Debates and Research Directions',
 'Integrative Case Conference','Chapter Synthesis','Applied Case Workshop','Chapter Summary'
]
for ch,d in chapters.items():
 d['objectives_clean']=parse_objectives(d['objectives']);clean=collections.OrderedDict()
 for heading,paras in d['sections'].items():
  out=[]
  for p in paras:
   if p.lower().startswith('verified chapter body word count:'):continue
   p=british_english(p)
   if p:out.append(p)
  if out:clean[heading]=out
 if list(clean)!=EXPECTED_SECTION_HEADINGS:raise ValueError(f'Chapter {ch} headings do not match the source sequence: {list(clean)}')
 d['sections_clean']=clean
for i,f in enumerate(FIGURES,1):
 f['number']=i;f['caption']=f'Figure {i}. {british_english(f["title"])}.';f['chapter_title']=CHAPTER_TITLES[i-1]

def make_diagram(fig):
 W,H=1600,860
 pal={'ink':'#243B53','muted':'#506579','accent':'#BD985B','card':'#FFFFFF','canvas':'#FFFFFF','center':'#EEF3F7','line':'#9AA9B7'}
 im=Image.new('RGB',(W,H),pal['canvas']);d=ImageDraw.Draw(im)
 header=ImageFont.truetype(SANS_BOLD,29);body=ImageFont.truetype(SANS,25);centerfont=ImageFont.truetype(SANS_BOLD,30)
 def wrap(text,font,width):
  lines=[];line=''
  for word in text.split():
   cand=word if not line else line+' '+word
   if d.textbbox((0,0),cand,font=font)[2]<=width:line=cand
   else:
    if line:lines.append(line)
    line=word
  if line:lines.append(line)
  return lines
 # Lines are behind filled cards, so they never strike through text.
 connectors=[((575,345),(545,205)),((1025,345),(1055,205)),((575,405),(545,420)),((1025,405),(1055,420)),((700,535),(510,650)),((900,535),(1090,650))]
 for a,b in connectors:d.line([a,b],fill=pal['line'],width=4)
 boxes=[(55,70,545,250),(1055,70,1545,250),(55,330,545,510),(1055,330,1545,510),(280,650,740,830),(860,650,1320,830)]
 for box,(label,content) in zip(boxes,fig['nodes']):
  x0,y0,x1,y1=box;d.rounded_rectangle(box,radius=22,fill=pal['card'],outline=pal['ink'],width=4)
  d.text((x0+24,y0+15),label,font=header,fill=pal['ink']);avail=x1-x0-48
  lines=wrap(content,body,avail);font=body;line_h=30;y=y0+55
  if y+len(lines)*line_h>y1-7:
   for size in (22,20,18,16):
    small=ImageFont.truetype(SANS,size);candidate=wrap(content,small,avail);candidate_y=y0+54;candidate_line_h=size+5
    if candidate_y+len(candidate)*candidate_line_h<=y1-9:
     lines=candidate;font=small;line_h=candidate_line_h;y=candidate_y;break
   else:raise ValueError(f'Figure node text does not fit: {label}: {content}')
  for line in lines:d.text((x0+24,y),line,font=font,fill=pal['muted']);y+=line_h
 cx0,cy0,cx1,cy1=575,290,1025,550
 d.rounded_rectangle((cx0,cy0,cx1,cy1),radius=28,fill=pal['center'],outline=pal['ink'],width=5)
 lines=wrap(fig['chapter_title'],centerfont,cx1-cx0-56);lh=38;y=cy0+((cy1-cy0)-len(lines)*lh)//2
 for line in lines:
  box=d.textbbox((0,0),line,font=centerfont);x=cx0+((cx1-cx0)-(box[2]-box[0]))//2
  d.text((x,y),line,font=centerfont,fill=pal['ink']);y+=lh
 out=BytesIO();im.save(out,format='PNG',optimize=True);out.seek(0);return out

def make_cover_image():
 im=Image.open(COVER_SOURCE).convert('RGB');w,h=im.size;target_h=round(w*10/7)
 if target_h>=h:
  canvas=Image.new('RGB',(w,target_h),(4,7,15));top=(target_h-h)//2;bottom=target_h-h-top
  if top:canvas.paste(im.crop((0,0,w,1)).resize((w,top)),(0,0))
  canvas.paste(im,(0,top))
  if bottom:canvas.paste(im.crop((0,h-1,w,h)).resize((w,bottom)),(0,top+h))
 else:
  target_w=round(h*7/10);left=(w-target_w)//2;canvas=im.crop((left,0,left+target_w,h))
 out=BytesIO();canvas.save(out,format='JPEG',quality=95,optimize=True);out.seek(0);return out

# DOCX utilities

def add_field(run,instruction,cached=''):
 r=run._r;begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
 instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=instruction
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
 text=OxmlElement('w:t');text.text=cached;end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
 for el in (begin,instr,sep,text,end):r.append(el)

def set_run_font(run,name='Times New Roman',size=11,bold=None,italic=None,color=None):
 run.font.name=name;run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name);run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 if italic is not None:run.italic=italic
 if color:run.font.color.rgb=RGBColor.from_string(color)

def set_doc_styles(doc):
 styles=doc.styles;normal=styles['Normal'];normal.font.name='Times New Roman';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');normal.font.size=Pt(10.75);normal.font.color.rgb=RGBColor(40,48,56)
 pf=normal.paragraph_format;pf.space_after=Pt(5);pf.line_spacing=1.12;pf.first_line_indent=Inches(.23);pf.widow_control=True
 for nm,size,col in [('Heading 1',23,'213B53'),('Heading 2',13.5,'284C68')]:
  st=styles[nm];st.font.name='Times New Roman';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(col);st.paragraph_format.keep_with_next=True;st.paragraph_format.widow_control=True
 styles['Heading 1'].paragraph_format.space_before=Pt(0);styles['Heading 1'].paragraph_format.space_after=Pt(12)
 styles['Heading 2'].paragraph_format.space_before=Pt(14);styles['Heading 2'].paragraph_format.space_after=Pt(6)
 if 'FrontMatter Heading' not in styles:
  from docx.enum.style import WD_STYLE_TYPE
  st=styles.add_style('FrontMatter Heading',WD_STYLE_TYPE.PARAGRAPH)
 else:st=styles['FrontMatter Heading']
 st.font.name='Times New Roman';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');st.font.size=Pt(20);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('213B53');st.paragraph_format.keep_with_next=True;st.paragraph_format.space_after=Pt(13);st.paragraph_format.first_line_indent=Inches(0)
 if 'Chapter Kicker' not in styles:
  from docx.enum.style import WD_STYLE_TYPE
  st=styles.add_style('Chapter Kicker',WD_STYLE_TYPE.PARAGRAPH)
 else:st=styles['Chapter Kicker']
 st.font.name='Times New Roman';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');st.font.size=Pt(10);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('BD985B');st.paragraph_format.keep_with_next=True;st.paragraph_format.space_after=Pt(5);st.paragraph_format.first_line_indent=Inches(0)
 styles['Caption'].font.name='Times New Roman';styles['Caption']._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');styles['Caption'].font.size=Pt(9.5);styles['Caption'].font.italic=True;styles['Caption'].font.color.rgb=RGBColor.from_string('506579')
 b=styles['List Bullet'];b.font.name='Times New Roman';b._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman');b.font.size=Pt(10.3);b.paragraph_format.left_indent=Inches(.25);b.paragraph_format.first_line_indent=Inches(-.16);b.paragraph_format.space_after=Pt(2);b.paragraph_format.line_spacing=1.08;b.paragraph_format.widow_control=True

def add_docx_toc_field(doc):
 p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(4)
 add_field(p.add_run(),'TOC \\o "1-2" \\h \\z \\u','Open in Word and update the table of contents to refresh page numbers.')
 upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings._element.append(upd)

def build_docx():
 doc=Document();set_doc_styles(doc);sec=doc.sections[0]
 sec.page_width=Inches(7);sec.page_height=Inches(10);sec.top_margin=Inches(0);sec.bottom_margin=Inches(0);sec.left_margin=Inches(0);sec.right_margin=Inches(0);sec.header_distance=Inches(0);sec.footer_distance=Inches(0)
 doc.core_properties.title='Advanced Psychopathology';doc.core_properties.subject='Doctoral-level textbook for counselling psychology';doc.core_properties.author='TRIUNE BIBLICAL UNIVERSITY GLOBAL EXTENSION, USA';doc.core_properties.keywords='CPY 902; psychopathology; diagnostic formulation; counselling psychology';doc.core_properties.category='Textbook';doc.core_properties.comments='Final edited manuscript. Author sign-off is required before publication release.'
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
 cover_run=p.add_run();set_run_font(cover_run,size=1);cover_shape=cover_run.add_picture(make_cover_image(),width=Inches(6.98));cover_shape._inline.docPr.set('descr','Front cover for Advanced Psychopathology (CPY 902), supplied by the author.');cover_shape._inline.docPr.set('title','Advanced Psychopathology front cover')
 sec=doc.add_section(WD_SECTION.NEW_PAGE)
 sec.page_width=Inches(7);sec.page_height=Inches(10);sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.78);sec.right_margin=Inches(.68);sec.header_distance=Inches(.34);sec.footer_distance=Inches(.34);sec.different_first_page_header_footer=True
 for part in (sec.header,sec.footer,sec.first_page_header,sec.first_page_footer):part.is_linked_to_previous=False
 h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.CENTER;h.paragraph_format.first_line_indent=Inches(0);h.paragraph_format.space_after=Pt(0);set_run_font(h.add_run('CPY. 902  |  ADVANCED PSYCHOPATHOLOGY'),size=8.5,color='506579')
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Inches(0);add_field(f.add_run(),'PAGE','1')
 sec.first_page_header.paragraphs[0].clear();sec.first_page_footer.paragraphs[0].clear()
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(80);set_run_font(p.add_run('TRIUNE BIBLICAL UNIVERSITY\nGLOBAL EXTENSION, USA'),size=17,bold=True,color='213B53')
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(35);set_run_font(p.add_run('CPY. 902'),size=16,bold=True,color='BD985B')
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(12);set_run_font(p.add_run('ADVANCED\nPSYCHOPATHOLOGY'),size=27,bold=True,color='213B53')
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(10);set_run_font(p.add_run('Doctoral-Level Textbook'),size=14,italic=True,color='506579')
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(25);set_run_font(p.add_run('TRIUNE BIBLICAL UNIVERSITY GLOBAL EXTENSION, USA'),size=11.5,bold=True)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);set_run_font(p.add_run('Shirley, New York, USA\nwww.tbuglobalextensionusa.com'),size=10.5)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(16);set_run_font(p.add_run('NOT FOR SALE'),size=12,bold=True,color='BD985B')
 doc.add_page_break()
 p=doc.add_paragraph('Publication and Academic Use Notice',style='Heading 1')
 for t in notice:doc.add_paragraph(british_english(t))
 doc.add_page_break();doc.add_paragraph('Preface',style='Heading 1')
 for t in preface:doc.add_paragraph(british_english(t))
 doc.add_page_break();doc.add_paragraph('Contents',style='FrontMatter Heading');add_docx_toc_field(doc)
 doc.add_page_break();doc.add_paragraph('List of Figures',style='Heading 1')
 for fd in FIGURES:
  p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.left_indent=Inches(.12);p.paragraph_format.space_after=Pt(3);set_run_font(p.add_run(fd['caption']),size=9.5,italic=True,color='506579')
 doc.add_page_break()
 for ch,d in chapters.items():
  p=doc.add_paragraph(style='Heading 1');p.paragraph_format.page_break_before=(ch>1);p.paragraph_format.first_line_indent=Inches(0);set_run_font(p.add_run(f'CHAPTER {ch} — {d["title"]}'),size=23,bold=True,color='213B53')
  p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(6);set_run_font(p.add_run('Learning objectives'),size=12,bold=True,color='284C68')
  for obj in d['objectives_clean']:doc.add_paragraph(obj,style='List Bullet')
  p=doc.add_paragraph(f'A visual overview of this chapter appears in Figure {ch}.');p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
  shape=p.add_run().add_picture(make_diagram(FIGURES[ch-1]),width=Inches(5.85));shape._inline.docPr.set('descr',FIGURES[ch-1]['caption']);shape._inline.docPr.set('title',FIGURES[ch-1]['caption'])
  p=doc.add_paragraph(FIGURES[ch-1]['caption'],style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(6)
  idx=0
  for heading,paras in d['sections_clean'].items():
   idx+=1;htext=heading
   p=doc.add_paragraph(style='Heading 2');p.paragraph_format.first_line_indent=Inches(0);set_run_font(p.add_run(f'{ch}.{idx} {htext}'),size=13.5,bold=True,color='284C68')
   for text in paras:doc.add_paragraph(british_english(text))
 doc.add_page_break();doc.add_paragraph('Selected References and Current Diagnostic Resources',style='Heading 1')
 refs=[
  [('American Psychiatric Association. (2022). ',False),('Diagnostic and statistical manual of mental disorders',True),(' (5th ed., text rev.; DSM-5-TR). American Psychiatric Association Publishing.',False)],
  [('American Psychiatric Association. (2025). ',False),('Updates to DSM-5-TR criteria and text: September 2025 update',True),('. American Psychiatric Association.',False)],
  [('World Health Organization. (2024). ',False),('Clinical descriptions and diagnostic requirements for ICD-11 mental, behavioural and neurodevelopmental disorders',True),('. World Health Organization.',False)],
  [('World Health Organization. (2022). ',False),('International classification of diseases, eleventh revision (ICD-11): Mortality and morbidity statistics',True),('. World Health Organization.',False)],
  [('National Institute of Mental Health. (n.d.). ',False),('Mental health information and statistics',True),('. National Institutes of Health.',False)],
  [('Substance Abuse and Mental Health Services Administration. (n.d.). ',False),('Evidence-based resources and mental health information',True),('. U.S. Department of Health and Human Services.',False)],
  [('Centers for Disease Control and Prevention. (n.d.). ',False),('Mental health and suicide prevention resources',True),('. U.S. Department of Health and Human Services.',False)],
 ]
 for ref in refs:
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.28);p.paragraph_format.first_line_indent=Inches(-.28);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.05
  for txt,ital in ref:set_run_font(p.add_run(txt),size=9.6,italic=ital)
 doc.save(DOCX_OUT);return DOCX_OUT

# PDF typesetting
pdfmetrics.registerFont(TTFont('SourceSerif4',FONT_REG));pdfmetrics.registerFont(TTFont('SourceSerif4-Bold',FONT_BOLD));pdfmetrics.registerFont(TTFont('SourceSerif4-Italic',FONT_ITALIC));pdfmetrics.registerFont(TTFont('SourceSerif4-BoldItalic',FONT_BOLDITALIC))
pdfmetrics.registerFontFamily('SourceSerif4',normal='SourceSerif4',bold='SourceSerif4-Bold',italic='SourceSerif4-Italic',boldItalic='SourceSerif4-BoldItalic')
class FilteredTOC(TableOfContents):
 def __init__(self,levels,**kwargs):self.allowed_levels=set(levels);super().__init__(**kwargs)
 def addEntry(self,level,text,pageNum,key=None):
  if level in self.allowed_levels:super().addEntry(level,text,pageNum,key)
class BookDoc(BaseDocTemplate):
 def beforeDocument(self):
  self.canv.setTitle('Advanced Psychopathology');self.canv.setAuthor('TRIUNE BIBLICAL UNIVERSITY GLOBAL EXTENSION, USA');self.canv.setSubject('Doctoral-level textbook for counselling psychology');self.canv.setKeywords('CPY 902, psychopathology, diagnostic formulation, counselling psychology')
 def afterFlowable(self,flowable):
  if not isinstance(flowable,Paragraph):return
  name=flowable.style.name;text=flowable.getPlainText()
  if name=='ChapterTitleTOC':
   m=re.match(r'CHAPTER\s+(\d+)',text);key='chapter-'+(m.group(1) if m else str(self.page));self.canv.bookmarkPage(key)
   try:self.canv.addOutlineEntry(text,key,0,False)
   except Exception:pass
   self.notify('TOCEntry',(0,text,self.page,key))
  elif name=='SectionTOC':
   key='section-'+str(self.page)+'-'+str(abs(hash(text)));self.canv.bookmarkPage(key)
   try:self.canv.addOutlineEntry(text,key,1,True)
   except Exception:pass
   self.notify('TOCEntry',(1,text,self.page,key))
  elif name=='FrontTOCEntry':
   key='front-'+str(abs(hash(text)));self.canv.bookmarkPage(key)
   try:self.canv.addOutlineEntry(text,key,0,False)
   except Exception:pass
   self.notify('TOCEntry',(0,text,self.page,key))
  elif name=='FigureTOC':
   m=re.match(r'Figure (\d+)',text);key='figure-'+(m.group(1) if m else str(self.page));self.canv.bookmarkPage(key)
   try:self.canv.addOutlineEntry(text,key,1,True)
   except Exception:pass
   self.notify('LOFEntry',(0,text,self.page,key))

def draw_body(canvas,doc):
 canvas.saveState();w,h=doc.pagesize
 if doc.page>1:
  canvas.setFont('SourceSerif4',8.3);canvas.setFillColor(colors.HexColor('#506579'));canvas.drawString(doc.leftMargin,h-.34*inch,'CPY. 902  |  ADVANCED PSYCHOPATHOLOGY')
  canvas.setStrokeColor(colors.HexColor('#CCD4DC'));canvas.setLineWidth(.45);canvas.line(doc.leftMargin,h-.42*inch,w-doc.rightMargin,h-.42*inch);canvas.line(doc.leftMargin,.42*inch,w-doc.rightMargin,.42*inch)
  canvas.setFont('SourceSerif4',9);canvas.drawCentredString(w/2,.23*inch,str(doc.page))
 canvas.restoreState()
def draw_title(canvas,doc):pass

def build_pdf():
 navy=colors.HexColor('#213B53');muted=colors.HexColor('#506579');gold=colors.HexColor('#BD985B')
 st=getSampleStyleSheet()
 st.add(ParagraphStyle(name='BodyBook',fontName='SourceSerif4',fontSize=10.2,leading=13.6,textColor=colors.HexColor('#26323C'),alignment=TA_JUSTIFY,firstLineIndent=18,spaceAfter=5.5,widowOrphanControl=1))
 st.add(ParagraphStyle(name='ChapterKicker',fontName='SourceSerif4-Bold',fontSize=10.5,leading=13,textColor=gold,alignment=TA_LEFT,spaceAfter=7,firstLineIndent=0))
 st.add(ParagraphStyle(name='ChapterTitleTOC',fontName='SourceSerif4-Bold',fontSize=23,leading=27,textColor=navy,alignment=TA_LEFT,spaceAfter=12,keepWithNext=True,firstLineIndent=0))
 st.add(ParagraphStyle(name='SectionTOC',fontName='SourceSerif4-Bold',fontSize=13.2,leading=16.2,textColor=navy,alignment=TA_LEFT,spaceBefore=12,spaceAfter=5,keepWithNext=True,firstLineIndent=0))
 st.add(ParagraphStyle(name='FrontHeading',fontName='SourceSerif4-Bold',fontSize=20,leading=24,textColor=navy,spaceAfter=13,keepWithNext=True,firstLineIndent=0))
 st.add(ParagraphStyle(name='FrontTOCEntry',fontName='SourceSerif4-Bold',fontSize=20,leading=24,textColor=navy,spaceAfter=13,keepWithNext=True,firstLineIndent=0))
 st.add(ParagraphStyle(name='SmallHeading',fontName='SourceSerif4-Bold',fontSize=12,leading=15,textColor=navy,spaceBefore=7,spaceAfter=4,keepWithNext=True,firstLineIndent=0))
 st.add(ParagraphStyle(name='Objective',fontName='SourceSerif4',fontSize=9.7,leading=12.1,textColor=colors.HexColor('#26323C'),leftIndent=14,firstLineIndent=-8,spaceAfter=2.2,bulletIndent=0,widowOrphanControl=1))
 st.add(ParagraphStyle(name='FigIntro',fontName='SourceSerif4-Italic',fontSize=9.2,leading=11.3,textColor=muted,spaceBefore=5,spaceAfter=4,firstLineIndent=0))
 st.add(ParagraphStyle(name='FigureTOC',fontName='SourceSerif4-Italic',fontSize=9.3,leading=11.5,textColor=muted,alignment=TA_LEFT,spaceBefore=4,spaceAfter=8,keepWithNext=False,firstLineIndent=0))
 st.add(ParagraphStyle(name='TOC0',fontName='SourceSerif4-Bold',fontSize=9.3,leading=12,textColor=navy,leftIndent=0,firstLineIndent=0,spaceBefore=3,spaceAfter=2,alignment=TA_LEFT))
 st.add(ParagraphStyle(name='TOC1',fontName='SourceSerif4',fontSize=8.6,leading=10.3,textColor=colors.HexColor('#394A5B'),leftIndent=18,firstLineIndent=0,spaceAfter=1,alignment=TA_LEFT))
 st.add(ParagraphStyle(name='LOFStyle',fontName='SourceSerif4',fontSize=9.4,leading=12,textColor=colors.HexColor('#394A5B'),spaceBefore=3,spaceAfter=4,alignment=TA_LEFT,firstLineIndent=0))
 st.add(ParagraphStyle(name='Ref',fontName='SourceSerif4',fontSize=9.3,leading=12.1,textColor=colors.HexColor('#26323C'),leftIndent=22,firstLineIndent=-22,spaceAfter=7,widowOrphanControl=1))
 st.add(ParagraphStyle(name='CoverUniversity',fontName='SourceSerif4-Bold',fontSize=15,leading=19,textColor=navy,alignment=TA_CENTER,spaceAfter=4,firstLineIndent=0))
 st.add(ParagraphStyle(name='CoverCourse',fontName='SourceSerif4-Bold',fontSize=15,leading=18,textColor=gold,alignment=TA_CENTER,spaceAfter=17,firstLineIndent=0))
 st.add(ParagraphStyle(name='CoverTitle',fontName='SourceSerif4-Bold',fontSize=29,leading=33,textColor=navy,alignment=TA_CENTER,spaceAfter=14,firstLineIndent=0))
 st.add(ParagraphStyle(name='CoverSub',fontName='SourceSerif4-Italic',fontSize=14,leading=18,textColor=muted,alignment=TA_CENTER,spaceAfter=20,firstLineIndent=0))
 st.add(ParagraphStyle(name='CoverInfo',fontName='SourceSerif4',fontSize=10.5,leading=14,textColor=colors.HexColor('#26323C'),alignment=TA_CENTER,spaceAfter=4,firstLineIndent=0))
 st.add(ParagraphStyle(name='CoverAuthor',fontName='SourceSerif4-Bold',fontSize=11,leading=14,textColor=colors.HexColor('#26323C'),alignment=TA_CENTER,spaceAfter=6,firstLineIndent=0))
 st.add(ParagraphStyle(name='ProofStamp',fontName='SourceSerif4-Bold',fontSize=8.7,leading=11,textColor=colors.HexColor('#9B3F3F'),alignment=TA_CENTER,spaceBefore=16,firstLineIndent=0))
 st.add(ParagraphStyle(name='NoticeSmall',fontName='SourceSerif4-Italic',fontSize=8.4,leading=11,textColor=muted,spaceBefore=9,spaceAfter=3,firstLineIndent=0))
 pagesize=(7*inch,10*inch);left=.76*inch;right=.66*inch;top=.76*inch;bottom=.65*inch
 doc=BookDoc(str(PDF_OUT),pagesize=pagesize,leftMargin=left,rightMargin=right,topMargin=top,bottomMargin=bottom,title='Advanced Psychopathology',author='TRIUNE BIBLICAL UNIVERSITY GLOBAL EXTENSION, USA',subject='Doctoral-level textbook for counselling psychology')
 frame=Frame(left,bottom,pagesize[0]-left-right,pagesize[1]-top-bottom,id='normal',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
 cover_frame=Frame(0,0,pagesize[0],pagesize[1],id='cover',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
 doc.addPageTemplates([PageTemplate(id='cover',frames=[cover_frame],onPage=draw_title),PageTemplate(id='title',frames=[frame],onPage=draw_title),PageTemplate(id='body',frames=[frame],onPage=draw_body)])
 toc=FilteredTOC(levels=[0,1],levelStyles=[st['TOC0'],st['TOC1']],dotsMinLevel=0,rightColumnWidth=.38*inch)
 lof=FilteredTOC(levels=[0],notifyKind='LOFEntry',levelStyles=[st['LOFStyle']],dotsMinLevel=0,rightColumnWidth=.38*inch)
 cover_img=RLImage(make_cover_image(),width=pagesize[0],height=pagesize[1]);cover_img.hAlign='CENTER'
 story=[cover_img,NextPageTemplate('title'),PageBreak(),Spacer(1,.95*inch),Paragraph('TRIUNE BIBLICAL UNIVERSITY<br/>GLOBAL EXTENSION, USA',st['CoverUniversity']),Spacer(1,.28*inch),Paragraph('CPY. 902',st['CoverCourse']),Paragraph('ADVANCED<br/>PSYCHOPATHOLOGY',st['CoverTitle']),Paragraph('Doctoral-Level Textbook',st['CoverSub']),Spacer(1,.20*inch),Paragraph('TRIUNE BIBLICAL UNIVERSITY GLOBAL EXTENSION, USA',st['CoverAuthor']),Paragraph('Shirley, New York, USA<br/>www.tbuglobalextensionusa.com',st['CoverInfo']),Spacer(1,.16*inch),Paragraph('NOT FOR SALE',st['CoverCourse']),NextPageTemplate('body'),PageBreak()]
 story.append(Paragraph('Publication and Academic Use Notice',st['FrontTOCEntry']))
 for t in notice:story.append(Paragraph(escape(british_english(t)),st['BodyBook']))
 story.extend([PageBreak(),Paragraph('Preface',st['FrontTOCEntry'])])
 for t in preface:story.append(Paragraph(escape(british_english(t)),st['BodyBook']))
 story.extend([PageBreak(),Paragraph('Contents',st['FrontHeading']),toc,PageBreak(),Paragraph('List of Figures',st['FrontTOCEntry']),lof,PageBreak()])
 for ch,d in chapters.items():
  if ch>1:story.append(PageBreak())
  title=f'CHAPTER {ch} — {british_english(d["title"])}'
  story.append(Paragraph(escape(title),st['ChapterTitleTOC']));story.append(Paragraph('Learning objectives',st['SmallHeading']))
  for obj in d['objectives_clean']:story.append(Paragraph('•&nbsp;&nbsp;'+escape(obj),st['Objective']))
  story.append(Paragraph(f'A visual overview of this chapter appears in Figure {ch}.',st['FigIntro']))
  pic=make_diagram(FIGURES[ch-1]);pic.seek(0)
  story.append(KeepTogether([RLImage(pic,width=5.45*inch,height=5.45*inch*860/1600),Paragraph(escape(FIGURES[ch-1]['caption']),st['FigureTOC'])]))
  idx=0
  for heading,paras in d['sections_clean'].items():
   idx+=1;htext=heading;sectitle=f'{ch}.{idx} {htext}'
   story.append(Paragraph(escape(sectitle),st['SectionTOC']))
   for text in paras:story.append(Paragraph(escape(british_english(text)),st['BodyBook']))
 story.extend([PageBreak(),Paragraph('Selected References and Current Diagnostic Resources',st['FrontTOCEntry'])])
 refs=[
  [('American Psychiatric Association. (2022). ',False),('Diagnostic and statistical manual of mental disorders',True),(' (5th ed., text rev.; DSM-5-TR). American Psychiatric Association Publishing.',False)],
  [('American Psychiatric Association. (2025). ',False),('Updates to DSM-5-TR criteria and text: September 2025 update',True),('. American Psychiatric Association.',False)],
  [('World Health Organization. (2024). ',False),('Clinical descriptions and diagnostic requirements for ICD-11 mental, behavioural and neurodevelopmental disorders',True),('. World Health Organization.',False)],
  [('World Health Organization. (2022). ',False),('International classification of diseases, eleventh revision (ICD-11): Mortality and morbidity statistics',True),('. World Health Organization.',False)],
  [('National Institute of Mental Health. (n.d.). ',False),('Mental health information and statistics',True),('. National Institutes of Health.',False)],
  [('Substance Abuse and Mental Health Services Administration. (n.d.). ',False),('Evidence-based resources and mental health information',True),('. U.S. Department of Health and Human Services.',False)],
  [('Centers for Disease Control and Prevention. (n.d.). ',False),('Mental health and suicide prevention resources',True),('. U.S. Department of Health and Human Services.',False)],
 ]
 for ref in refs:
  text=''.join(f'<i>{escape(t)}</i>' if italic else escape(t) for t,italic in ref).replace('[[BR]]','<br/>');story.append(Paragraph(text,st['Ref']))
 doc.multiBuild(story);return PDF_OUT,toc,lof

print('PARSED CHAPTERS')
for ch,d in chapters.items():
 wc=sum(len(re.findall(r"\b[\w’'-]+\b",p)) for arr in d['sections_clean'].values() for p in arr)
 print(ch,repr(d['title']),'objectives',len(d['objectives_clean']),'sections',len(d['sections_clean']),'words',wc,'source_claim',d['source_word_count'],'difference',wc-(d['source_word_count'] or 0))
print('Front matter source paragraphs:',{k:len(v) for k,v in front.items()})
print('Generating editable DOCX...');build_docx()
print('Generating typeset PDF...');out,toc,lof=build_pdf()
print('Output sizes',out.stat().st_size,DOCX_OUT.stat().st_size)
final=pymupdf.open(str(out));full='\n'.join(p.get_text() for p in final)
print('PDF pages',final.page_count,'outline entries',len(final.get_toc()),'text words',len(re.findall(r"\b[\w’'-]+\b",full)))
print('Unresolved standard placeholders in manuscript',len(re.findall(r'\[[^\]]{1,120}\]|\b(?:TBD|TODO|TBA|XXX|Lorem ipsum|SOURCE NEEDED)\b',full,re.I)))
print('Religious/devotional terms excluding institutional name',{t:len(re.findall(t,full,re.I)) for t in ['prayer','devotional','scripture','worship','sermon','theological','Christianity','Islam']})
