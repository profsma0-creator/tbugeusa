from pathlib import Path
import re
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
base=Path('deliverables/Editorial_Assessment_and_Author_Queries')
text=base.with_suffix('.md').read_text()
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BookBody',fontName='Times-Roman',fontSize=11,leading=14,spaceAfter=7))
for name,size in [('Title',25),('Heading1',18),('Heading2',15),('Heading3',12)]:
 styles[name].fontName='Times-Bold';styles[name].fontSize=size;styles[name].leading=size+4;styles[name].textColor=colors.HexColor('#18384a')
 styles[name].spaceBefore=12;styles[name].spaceAfter=8
styles.add(ParagraphStyle(name='ItemBody',parent=styles['BookBody'],leftIndent=12,firstLineIndent=-7))
def inline(s):
 s=escape(s)
 s=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',s)
 s=re.sub(r'`(.*?)`',r'\1',s)
 return s
blocks=re.split(r'\n\s*\n',text)
story=[];doc=Document();sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.75)
sec.left_margin=sec.right_margin=Inches(.8)
doc.styles['Normal'].font.name='Cambria';doc.styles['Normal'].font.size=Pt(11)
for block in blocks:
 lines=block.splitlines()
 for line in lines if all(re.match(r'^(?:[-#]|\d+\.)',l) for l in lines) else [' '.join(lines)]:
  m=re.match(r'^(#{1,4}) (.*)',line)
  if m:
   level=len(m[1]);s=m[2];style=['Title','Heading1','Heading2','Heading3'][level-1]
   story.append(Paragraph(inline(s),styles[style]));doc.add_heading(s,level=0 if level==1 else level-1)
  else:
   item=bool(re.match(r'^(?:- |\d+\. )',line))
   s=re.sub(r'^- ','• ',line)
   story.append(Paragraph(inline(s),styles['ItemBody' if item else 'BookBody']))
   para=doc.add_paragraph()
   for j,part in enumerate(re.split(r'\*\*(.*?)\*\*',re.sub(r'`([^`]+)`',r'\1',s))):
    r=para.add_run(part);r.bold=bool(j%2)
header=sec.header.paragraphs[0];header.text='TRT 903  |  PRELIMINARY EDITORIAL ASSESSMENT'
header.style=doc.styles['Caption']
f=sec.footer.paragraphs[0];f.text='Not an edited manuscript • '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
doc.save(base.with_suffix('.docx'))
def furniture(c,d):
 c.saveState();w,h=d.pagesize;c.setStrokeColor(colors.HexColor('#c4ccd1'));c.line(50,h-38,w-50,h-38)
 c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#425868'))
 c.drawString(50,h-28,'TRT 903  |  PRELIMINARY EDITORIAL ASSESSMENT')
 c.drawString(50,28,'Not an edited manuscript • 21 October 2026')
 c.drawRightString(w-50,28,str(d.page));c.restoreState()
SimpleDocTemplate(str(base.with_suffix('.pdf')),pagesize=(504,720),leftMargin=50,rightMargin=50,topMargin=52,bottomMargin=48,title='Advanced Neurobiology of Trauma — Preliminary Editorial Assessment',author='').build(story,onFirstPage=furniture,onLaterPages=furniture)
